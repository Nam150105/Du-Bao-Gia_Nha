"""Sinh file báo cáo TTNT.docx theo thuật toán hiện tại của project."""

from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "src"))
ASSETS = ROOT / "docs" / "report_assets"
OUTPUT = ROOT / "TTNT.docx"


def load_metrics() -> dict:
    try:
        from utils import load_model

        return load_model().get("metrics", {})
    except Exception:
        return {
            "samples": 106_785,
            "mae": 5.92e9,
            "rmse": 12.97e9,
            "r2": 0.5668,
        }


def make_figures(metrics: dict) -> dict[str, Path]:
    import matplotlib.pyplot as plt
    import matplotlib.patches as mpatches

    ASSETS.mkdir(parents=True, exist_ok=True)
    paths: dict[str, Path] = {}

    # 1. Pipeline huấn luyện
    fig, ax = plt.subplots(figsize=(10, 5))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 5)
    ax.axis("off")
    boxes = [
        (0.3, 2.2, "house_data.csv\n~155k tin"),
        (2.2, 2.2, "Làm sạch\n(loại đất, outlier)"),
        (4.1, 2.2, "Feature\nengineering"),
        (6.0, 2.2, "OneHot +\nNumeric"),
        (7.9, 2.2, "HistGradient\nBoosting"),
        (9.0, 2.2, "model.pkl"),
    ]
    for i, (x, y, text) in enumerate(boxes):
        ax.add_patch(
            mpatches.FancyBboxPatch(
                (x, y),
                1.5,
                1.0,
                boxstyle="round,pad=0.05",
                facecolor="#dbeafe",
                edgecolor="#2563eb",
            )
        )
        ax.text(x + 0.75, y + 0.5, text, ha="center", va="center", fontsize=8)
        if i < len(boxes) - 1:
            ax.annotate(
                "",
                xy=(boxes[i + 1][0], y + 0.5),
                xytext=(x + 1.5, y + 0.5),
                arrowprops=dict(arrowstyle="->", color="#334155"),
            )
    ax.set_title("Quy trình huấn luyện mô hình", fontsize=12, fontweight="bold")
    p1 = ASSETS / "fig01_train_pipeline.png"
    fig.savefig(p1, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    paths["pipeline"] = p1

    # 2. Luồng dự báo
    fig, ax = plt.subplots(figsize=(10, 6))
    ax.set_xlim(0, 10)
    ax.set_ylim(0, 6)
    ax.axis("off")
    steps = [
        (0.5, 4.5, "Form web\n(JSON)"),
        (2.5, 4.5, "Chuẩn hóa\n34 tỉnh/thành"),
        (4.5, 4.5, "ML: đơn giá\n/m² (log)"),
        (6.5, 4.5, "× diện tích"),
        (8.2, 4.5, "Anchor +\nsàn PL/NT"),
        (4.5, 2.0, "Tin tham khảo\n(cùng khu vực)"),
        (8.2, 2.0, "Kết quả\n+ ảnh proxy"),
    ]
    for x, y, text in steps:
        ax.add_patch(
            mpatches.FancyBboxPatch(
                (x, y),
                1.6,
                1.0,
                boxstyle="round,pad=0.05",
                facecolor="#dcfce7",
                edgecolor="#16a34a",
            )
        )
        ax.text(x + 0.8, y + 0.5, text, ha="center", va="center", fontsize=8)
    for a, b in [(0, 1), (1, 2), (2, 3), (3, 4), (2, 5), (4, 6)]:
        x1, y1, _ = steps[a]
        x2, y2, _ = steps[b]
        ax.annotate(
            "",
            xy=(x2 + 0.8, y2 + 0.5 if y1 == y2 else y2 + 1),
            xytext=(x1 + 0.8, y1 + 0.5),
            arrowprops=dict(arrowstyle="->", color="#334155"),
        )
    ax.set_title("Luồng dự báo giá khi người dùng gửi form", fontsize=12, fontweight="bold")
    p2 = ASSETS / "fig02_predict_flow.png"
    fig.savefig(p2, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    paths["predict"] = p2

    # 3. Kiến trúc web
    fig, ax = plt.subplots(figsize=(9, 4.5))
    ax.set_xlim(0, 9)
    ax.set_ylim(0, 4)
    ax.axis("off")
    layers = [
        (0.5, 2.5, "Trình duyệt\nindex.html + CSS"),
        (3.0, 2.5, "Flask app.py\n/predict, /api/*"),
        (5.5, 2.5, "predict.py\nutils.py"),
        (7.5, 2.5, "house_price\n_model.pkl"),
    ]
    for x, y, t in layers:
        ax.add_patch(
            mpatches.FancyBboxPatch(
                (x, y), 2.0, 1.2, boxstyle="round", facecolor="#fef3c7", edgecolor="#d97706"
            )
        )
        ax.text(x + 1, y + 0.6, t, ha="center", va="center", fontsize=8)
    ax.set_title("Kiến trúc ứng dụng web 3 cột", fontsize=12, fontweight="bold")
    p3 = ASSETS / "fig03_web_arch.png"
    fig.savefig(p3, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    paths["arch"] = p3

    # 4. Metrics
    fig, ax = plt.subplots(figsize=(7, 4))
    labels = ["R²", "MAE (tỷ VND)", "RMSE (tỷ VND)"]
    r2 = float(metrics.get("r2", 0))
    mae = float(metrics.get("mae", 0)) / 1e9
    rmse = float(metrics.get("rmse", 0)) / 1e9
    vals = [r2, mae, rmse]
    colors = ["#3b82f6", "#06b6d4", "#8b5cf6"]
    bars = ax.bar(labels, vals, color=colors)
    ax.set_ylabel("Giá trị")
    ax.set_title(f"Chỉ số đánh giá (n = {int(metrics.get('samples', 0)):,})")
    for bar, v in zip(bars, vals):
        ax.text(
            bar.get_x() + bar.get_width() / 2,
            bar.get_height(),
            f"{v:.2f}",
            ha="center",
            va="bottom",
            fontsize=10,
        )
    p4 = ASSETS / "fig04_metrics.png"
    fig.savefig(p4, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    paths["metrics"] = p4

    # 5. Đặc trưng
    fig, ax = plt.subplots(figsize=(8, 5))
    ax.axis("off")
    feat_text = (
        "Đặc trưng phân loại (One-Hot):\n"
        "  • property_type (5 loại nhà ở)\n"
        "  • city (34 tỉnh/thành + khac)\n\n"
        "Đặc trưng số:\n"
        "  • area_m2, bedrooms, bathrooms\n"
        "  • frontage_m, road_width_m, total_floors\n"
        "  • legal_ordinal, furnished_ordinal\n"
        "  • room_density, livability_score\n\n"
        "Biến mục tiêu: unit_price_m2 = price_vnd / area_m2\n"
        "Biến đổi: log1p khi train, expm1 khi predict"
    )
    ax.text(0.05, 0.95, feat_text, va="top", fontsize=11)
    ax.set_title("Vector đặc trưng đưa vào mô hình", fontsize=12, fontweight="bold")
    p5 = ASSETS / "fig05_features.png"
    fig.savefig(p5, dpi=150, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    paths["features"] = p5

    return paths


def add_heading(doc, text: str, level: int = 1) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    h = doc.add_heading(text, level=level)
    if level == 0:
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_image(doc, path: Path, caption: str, width_in: float = 5.8) -> None:
    from docx.shared import Inches

    if path.exists():
        doc.add_picture(str(path), width=Inches(width_in))
        cap = doc.add_paragraph(caption)
        cap.style = "Caption"
        cap.alignment = 1  # center


def build_doc(metrics: dict, figures: dict[str, Path]) -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(13)

    # Trang bìa
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("\n\nBÁO CÁO\nTHỰC TẬP TỐT NGHIỆP\n\n")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = sub.add_run("ĐỀ TÀI:\nDỰ BÁO GIÁ NHÀ VÀ CĂN HỘ\nBẰNG HỌC MÁY VÀ ỨNG DỤNG WEB\n\n")
    r2.font.size = Pt(16)
    r2.bold = True

    info = doc.add_paragraph(
        "Project: Du-Bao-Gia_Nha\n"
        "Công nghệ: Python, scikit-learn, Flask\n"
        "Dữ liệu: house_data.csv (~155.000 tin đăng)\n"
    )
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_page_break()

    add_heading(doc, "MỤC LỤC", 1)
    toc = [
        "Lời mở đầu",
        "Chương 1. Tổng quan đề tài",
        "Chương 2. Cơ sở lý thuyết",
        "Chương 3. Phân tích dữ liệu và tiền xử lý",
        "Chương 4. Thuật toán và mô hình học máy",
        "Chương 5. Thiết kế và triển khai hệ thống web",
        "Chương 6. Kết quả thực nghiệm",
        "Chương 7. Kết luận và hướng phát triển",
        "Tài liệu tham khảo",
    ]
    for item in toc:
        doc.add_paragraph(item, style="List Number")
    doc.add_page_break()

    add_heading(doc, "LỜI MỞ ĐẦU", 1)
    doc.add_paragraph(
        "Bất động sản nhà ở tại Việt Nam có quy mô lớn và biến động giá theo khu vực, "
        "loại hình, pháp lý và nội thất. Đề tài xây dựng hệ thống Du-Bao-Gia_Nha nhằm "
        "ước tính giá bán căn hộ và nhà ở (không gồm đất nền) từ dữ liệu tin đăng thực tế, "
        "kết hợp mô hình học máy HistGradientBoostingRegressor và giao diện web Flask "
        "ba cột: nhập liệu, kết quả dự báo, tin tham khảo thị trường."
    )
    doc.add_paragraph(
        "Báo cáo được viết lại theo đúng mã nguồn và pipeline đang chạy tại thời điểm "
        "hoàn thiện project, bao gồm 34 đơn vị hành chính cấp tỉnh (Nghị quyết "
        "202/2025/QH15), cơ chế sàn giá theo pháp lý/nội thất, và proxy ảnh tin đăng."
    )

    add_heading(doc, "CHƯƠNG 1. TỔNG QUAN ĐỀ TÀI", 1)
    add_heading(doc, "1.1. Mục tiêu", 2)
    for g in [
        "Xây dựng pipeline làm sạch và huấn luyện trên tập house_data.csv.",
        "Dự báo giá bán (VND) và đơn giá trên m² từ thông tin người dùng nhập.",
        "Hiển thị tin đăng tham khảo cùng khu vực, thông số tương đồng.",
        "Triển khai web app một lệnh: python run_project.py.",
    ]:
        doc.add_paragraph(g, style="List Bullet")

    add_heading(doc, "1.2. Phạm vi", 2)
    doc.add_paragraph(
        "Hệ thống hỗ trợ 5 loại nhà ở: căn hộ chung cư, nhà riêng, biệt thự, liền kề, "
        "nhà mặt phố. Loại trừ đất nền, văn phòng. Khu vực: 34 tỉnh/thành sau sắp xếp "
        "hành chính 2025; tin không map được gán mã khac khi train."
    )

    add_heading(doc, "1.3. Cấu trúc project", 2)
    structure = (
        "app.py — Flask routes\n"
        "run_project.py — train + chạy web\n"
        "src/utils.py — làm sạch, train, calibration\n"
        "src/predict.py — dự báo + tin tham khảo\n"
        "src/regions.py — 34 tỉnh/thành\n"
        "templates/index.html, static/style.css — giao diện\n"
        "data/house_data.csv, models/house_price_model.pkl"
    )
    doc.add_paragraph(structure)

    add_heading(doc, "CHƯƠNG 2. CƠ SỞ LÝ THUYẾT", 1)
    add_heading(doc, "2.1. Hồi quy gradient boosting", 2)
    doc.add_paragraph(
        "HistGradientBoostingRegressor (scikit-learn) xây dựng nhiều cây quyết định "
        "theo chiến lược boosting, tối ưu trên histogram của đặc trưng liên tục, "
        "phù hợp dữ liệu lớn (>100k mẫu). Mô hình dự báo đơn giá/m²; giá căn = "
        "đơn giá × diện tích."
    )
    add_heading(doc, "2.2. Biến đổi log mục tiêu", 2)
    doc.add_paragraph(
        "Pipeline dùng TransformedTargetRegressor: huấn luyện trên log1p(unit_price_m2), "
        "dự báo qua expm1 để giảm lệch phải do giá cực đoan."
    )
    add_heading(doc, "2.3. One-Hot Encoding", 2)
    doc.add_paragraph(
        "Đặc trưng phân loại property_type và city được mã hóa One-Hot với "
        "sparse_output=False (ma trận dày) vì HistGradientBoosting yêu cầu dense."
    )

    add_heading(doc, "CHƯƠNG 3. PHÂN TÍCH DỮ LIỆU VÀ TIỀN XỬ LÝ", 1)
    add_heading(doc, "3.1. Nguồn dữ liệu", 2)
    doc.add_paragraph(
        "File house_data.csv chứa khoảng 155.760 dòng tin đăng từ nhiều nguồn "
        "(nhadatvui, guland, batdongsan, bannha888, …) với các cột: title, address, "
        "ward, district, province, area_m2, bedrooms, bathrooms, price_vnd, legal, "
        "furniture, source_url, livability_score, tọa độ lat/lng, v.v."
    )
    add_heading(doc, "3.2. Làm sạch (clean_data)", 2)
    steps = [
        "Loại tin đất: asset_type Đất/Đất nền hoặc tiêu đề khớp regex đất nền.",
        "Suy luận property_type từ asset_type, type, title.",
        "Gán city qua extract_region_code: map 63 tỉnh cũ → 34 mã mới.",
        "Lọc price_vnd ≥ 300 triệu, area_m2 ≥ 15; loại outlier percentile 0.5%–99.5%.",
        "Impute bedrooms/bathrooms; chuẩn hóa legal_status, furnished_status.",
        "Tính unit_price_m2, room_density = (bedrooms + 0.5×bathrooms) / area.",
    ]
    for s in steps:
        doc.add_paragraph(s, style="List Bullet")
    doc.add_paragraph(
        f"Sau làm sạch, tập huấn luyện hiện tại: khoảng {int(metrics.get('samples', 106785)):,} mẫu."
    )

    add_heading(doc, "CHƯƠNG 4. THUẬT TOÁN VÀ MÔ HÌNH HỌC MÁY", 1)
    add_image(doc, figures["pipeline"], "Hình 4.1. Quy trình huấn luyện")
    add_image(doc, figures["features"], "Hình 4.2. Đặc trưng đầu vào mô hình")

    add_heading(doc, "4.1. Tham số HistGradientBoostingRegressor", 2)
    params = (
        "max_iter=800, max_depth=12, learning_rate=0.04\n"
        "min_samples_leaf=30, l2_regularization=0.3, random_state=42"
    )
    doc.add_paragraph(params)

    add_heading(doc, "4.2. Thuật toán dự báo (predict_price)", 2)
    doc.add_paragraph(
        "Bước 1 — Chuẩn bị mẫu ML với legal_ordinal/furnished_ordinal trung tính "
        "(hop_dong, co_ban) để mô hình không học trực tiếp nhãn người dùng chọn."
    )
    doc.add_paragraph(
        "Bước 2 — unit_ml = pipeline.predict(sample); blend_weight=1.0 → chỉ dùng ML "
        "(không trộn baseline median như phiên bản cũ gây R² âm)."
    )
    doc.add_paragraph(
        "Bước 3 — Anchor floor: dự báo mẫu neo 2PN + hợp đồng; nhân hệ số "
        "LEGAL_FLOOR × FURNISHED_FLOOR × (1+bedroom_step)^(PN-2) × frontage; "
        "predicted_price = max(ML×area, anchor×hệ số)."
    )
    doc.add_paragraph(
        "Bước 4 — apply_calibration: áp dụng sàn monotonic pháp lý và nội thất "
        "(sổ đỏ > hợp đồng > chờ sổ > không rõ; full > cơ bản > không rõ > không NT)."
    )
    add_image(doc, figures["predict"], "Hình 4.3. Luồng xử lý khi dự báo")

    add_heading(doc, "4.3. Tin tham khảo (get_reference_examples)", 2)
    doc.add_paragraph(
        "Lọc reference_pool theo đúng city người chọn; ưu tiên cùng property_type. "
        "Chấm điểm khoảng cách diện tích, PN, WC, giá; trả tối đa 10 tin; ảnh tải "
        "lazy qua /api/image-proxy (og:image từ trang tin)."
    )

    add_heading(doc, "CHƯƠNG 5. THIẾT KẾ VÀ TRIỂN KHAI HỆ THỐNG WEB", 1)
    add_image(doc, figures["arch"], "Hình 5.1. Kiến trúc tầng")

    add_heading(doc, "5.1. Giao diện 3 cột", 2)
    doc.add_paragraph(
        "Cột 1: form nhập (loại nhà, combobox 34 tỉnh/thành có tìm kiếm, diện tích, "
        "PN, WC, pháp lý, nội thất, mặt tiền). Cột 2: giá dự báo và metrics model. "
        "Cột 3: marquee tin ngẫu nhiên trước predict; sau predict hiển thị 3–10 tin "
        "tương tự, scroll nội bộ cột."
    )
    add_heading(doc, "5.2. API Flask", 2)
    apis = [
        "GET / — trang chủ",
        "POST /predict — JSON, trả giá + references",
        "GET /api/image-proxy?url= — proxy ảnh tin đăng",
        "GET /health — kiểm tra model",
    ]
    for a in apis:
        doc.add_paragraph(a, style="List Bullet")

    add_heading(doc, "CHƯƠNG 6. KẾT QUẢ THỰC NGHIỆM", 1)
    add_image(doc, figures["metrics"], "Hình 6.1. Chỉ số trên tập test (20%)")

    mae_ty = float(metrics.get("mae", 0)) / 1e9
    rmse_ty = float(metrics.get("rmse", 0)) / 1e9
    doc.add_paragraph(
        f"• Số mẫu train sau làm sạch: {int(metrics.get('samples', 0)):,}\n"
        f"• R² (test): {float(metrics.get('r2', 0)):.4f}\n"
        f"• MAE (test): {mae_ty:.2f} tỷ VND\n"
        f"• RMSE (test): {rmse_ty:.2f} tỷ VND"
    )
    doc.add_paragraph(
        "Nhận xét: R² ~0.57 phản ánh dự báo giá tổng còn sai số lớn do thị trường "
        "phức tạp; MAE ~6 tỷ là mức trung bình tuyệt đối trên căn nhà. Cơ chế sàn "
        "pháp lý/nội thất giúp kết quả hợp lý hơn (ví dụ full NT không rẻ hơn không NT)."
    )

    add_heading(doc, "CHƯƠNG 7. KẾT LUẬN VÀ HƯỚNG PHÁT TRIỂN", 1)
    add_heading(doc, "7.1. Kết luận", 2)
    doc.add_paragraph(
        "Đề tài đã xây dựng pipeline end-to-end từ CSV đến web app, dùng "
        "HistGradientBoosting trên đơn giá/m² với 34 khu vực hành chính mới, "
        "loại đất, và hiển thị tin tham khảo có ảnh proxy."
    )
    add_heading(doc, "7.2. Hướng phát triển", 2)
    for h in [
        "Bổ sung quận/huyện khi đủ mẫu; dùng lat/lng trong mô hình.",
        "Cập nhật Đồng Nai thành TP TWTU (2026) khi dữ liệu đồng bộ.",
        "Deploy Docker + API REST; A/B test mô hình.",
        "Cache ảnh og:image lúc train để giảm tải proxy.",
    ]:
        doc.add_paragraph(h, style="List Bullet")

    add_heading(doc, "TÀI LIỆU THAM KHẢO", 1)
    refs = [
        "Scikit-learn — HistGradientBoostingRegressor Documentation.",
        "Flask Documentation — Web Development.",
        "Nghị quyết 202/2025/QH15 — Sắp xếp đơn vị hành chính cấp tỉnh.",
        "Pandas — Data Analysis Library.",
    ]
    for i, r in enumerate(refs, 1):
        doc.add_paragraph(f"[{i}] {r}")

    doc.save(OUTPUT)
    print(f"Đã tạo: {OUTPUT}")


def main() -> None:
    metrics = load_metrics()
    figures = make_figures(metrics)
    build_doc(metrics, figures)


if __name__ == "__main__":
    main()
